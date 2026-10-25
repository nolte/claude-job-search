"""Export tests for skills/job-application-export/scripts/export_document.py (spec CG-12).

Uses a synthetic Lebenslauf. Requires pandoc for md/html/docx and python-docx for the DOCX
structure check. The PDF export needs LuaLaTeX; without it the PDF test is skipped and the
skip reason is printed, so the result is not reported as passed.
"""

import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXPORTER = os.path.join(ROOT, "skills", "job-application-export", "scripts", "export_document.py")

SOURCE = """\
---
grundlage_hash: 0000000000000000000000000000000000000000000000000000000000000000
format: chronologisch
---

# Lebenslauf

## Persönliche Daten
Max Mustermann · Musterstraße 1 · 00000 Musterstadt

## Berufserfahrung

### 04/2021 – heute | Senior Backend-Entwickler | Musterfirma GmbH, Berlin
- Betreuung der Event-Pipeline in Python und Kafka; Latenz von 800 ms auf 120 ms gesenkt.
- Einführung von Code-Reviews im Team mit fünf Entwicklern.

## Sprachkenntnisse (GER-Niveau)
- Deutsch: Muttersprache

Ort, Datum
"""

HEADINGS = ["Persönliche Daten", "Berufserfahrung", "Sprachkenntnisse (GER-Niveau)"]


def have(tool):
    return shutil.which(tool) is not None


def python_docx_available():
    try:
        import docx  # noqa: F401
    except ImportError:
        return False
    return True


class ToolchainTests(unittest.TestCase):
    def test_ci_runs_the_tests_that_otherwise_skip(self):
        # F-02: without pandoc, python-docx and PyYAML the export and frontmatter tests skip and the job stays green.
        # LuaLaTeX (the PDF tests) is not required: the PDF checks run locally (docs/en/evals.md, "What the CI covers").
        if not os.environ.get("CI"):
            self.skipTest("not in CI: missing tools only skip the dependent tests")
        try:
            import yaml  # noqa: F401
        except ImportError:
            yaml = None
        missing = [name for name, ok in (("pandoc", have("pandoc")), ("python-docx", python_docx_available()),
                                          ("PyYAML", yaml is not None)) if not ok]
        self.assertEqual(missing, [], "install them in the test job (tests/requirements.txt, apt pandoc)")


@unittest.skipUnless(have("pandoc"), "pandoc is not installed")
class ExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.source = os.path.join(self.tmp, "lebenslauf.md")
        with open(self.source, "w", encoding="utf-8") as handle:
            handle.write(SOURCE)
        self.out = os.path.join(self.tmp, "out")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def export(self, fmt):
        result = subprocess.run(
            [sys.executable, EXPORTER, self.source, "--format", fmt, "--doc-type", "cv", "--out-dir", self.out, "--name", "Test_Lebenslauf"],
            capture_output=True,
            text=True,
        )
        return result

    def test_markdown_export_keeps_content(self):
        # Scenario J: export is a copy of the approved content
        result = self.export("md")
        self.assertEqual(result.returncode, 0, result.stderr)
        output = os.path.join(self.out, "Test_Lebenslauf.md")
        with open(output, encoding="utf-8") as handle:
            text = handle.read()
        for line in SOURCE.splitlines():
            if line.startswith("- "):
                self.assertIn(line[2:], text)

    def test_html_export_passes_content_check(self):
        result = self.export("html")
        self.assertEqual(result.returncode, 0, result.stderr)
        with open(os.path.join(self.out, "Test_Lebenslauf.html"), encoding="utf-8") as handle:
            self.assertIn("Latenz von 800 ms", handle.read())

    def test_motivation_doc_type_is_accepted_with_a_two_page_limit(self):
        # MS-7: the Motivationsschreiben has its own document type; pandoc only, no LuaLaTeX
        sys.path.insert(0, os.path.dirname(EXPORTER))
        import export_document

        self.assertEqual(export_document.PAGE_LIMITS["motivation"], 2)
        result = subprocess.run(
            [sys.executable, EXPORTER, self.source, "--format", "tex", "--doc-type", "motivation", "--out-dir", self.out, "--name", "Test_Motivation"],
            capture_output=True,
            text=True,
        )
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_open_placeholders_are_reported_and_stay_visible_without_blocking(self):
        # audit finding EXP-001: the export used to say "Inhaltsprüfung: ok" and nothing else
        with open(self.source, "a", encoding="utf-8") as handle:
            handle.write("\n- [ERGÄNZEN: Aufgaben]\n- [KLÄREN: Titel laut Zeugnis]\n")
        for fmt in ("md", "html"):
            result = self.export(fmt)
            self.assertEqual(result.returncode, 0, f"{fmt}: {result.stderr}")
            self.assertIn("Platzhalter: 2 offen", result.stdout, fmt)
            self.assertIn("2 offene Platzhalter", result.stderr, fmt)
        with open(os.path.join(self.out, "Test_Lebenslauf.html"), encoding="utf-8") as handle:
            self.assertIn("ERGÄNZEN", handle.read())

    def test_a_placeholder_in_the_frontmatter_is_counted_like_the_gate_does(self):
        with open(self.source, encoding="utf-8") as handle:
            text = handle.read()
        with open(self.source, "w", encoding="utf-8") as handle:
            handle.write(text.replace("format: chronologisch", "format: chronologisch\nnotiz: [ERGÄNZEN: Ort]", 1))
        result = self.export("md")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Platzhalter: 1 offen", result.stdout)

    def test_a_document_without_placeholders_prints_no_placeholder_line(self):
        result = self.export("md")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertNotIn("Platzhalter", result.stdout)
        self.assertNotIn("Platzhalter", result.stderr)

    def test_email_doc_type_only_exports_markdown(self):
        # R-10.2: the e-mail text exists only as Markdown
        for fmt in ("pdf", "docx", "html", "tex"):
            result = subprocess.run(
                [sys.executable, EXPORTER, self.source, "--format", fmt, "--doc-type", "email", "--out-dir", self.out, "--name", "Test_Email"],
                capture_output=True,
                text=True,
            )
            self.assertEqual(result.returncode, 1, fmt)
            self.assertIn("nur als Markdown", result.stderr)
        result = subprocess.run(
            [sys.executable, EXPORTER, self.source, "--format", "md", "--doc-type", "email", "--out-dir", self.out, "--name", "Test_Email"],
            capture_output=True,
            text=True,
        )
        self.assertEqual(result.returncode, 0, result.stderr)

    def run_main_with_pages(self, doc_type, pages, source_text=None):
        # the PDF step is replaced, so the page-limit path runs without LuaLaTeX
        from unittest import mock
        import contextlib
        import io

        sys.path.insert(0, os.path.dirname(EXPORTER))
        import export_document

        if source_text is not None:
            with open(self.source, "w", encoding="utf-8") as handle:
                handle.write(source_text)

        def fake_pdf(pandoc, source, out_dir, name, grundlage_hash, doc_type="cv", epoch=None):
            (out_dir / f"{name}.pdf").write_bytes(b"%PDF-fake")
            text = export_document.read_document(source)[1]
            return text, pages

        err = io.StringIO()
        with mock.patch.object(export_document, "export_pdf", fake_pdf), contextlib.redirect_stdout(io.StringIO()), \
                contextlib.redirect_stderr(err):
            code = export_document.main([self.source, "--format", "pdf", "--doc-type", doc_type, "--out-dir", self.out, "--name", "Test_Doc"])
        return code, err.getvalue()

    def test_page_limit_overrun_has_its_own_exit_code_and_keeps_file_and_manifest(self):
        # LB-2.3, AS-7.1: 4 = over the page limit; file and manifest stay
        code, err = self.run_main_with_pages("letter", 2)
        self.assertEqual(code, 4)
        self.assertIn("Seitenlimit überschritten", err)
        self.assertTrue(os.path.exists(os.path.join(self.out, "Test_Doc.pdf")))
        with open(os.path.join(self.out, "Test_Doc.pdf.manifest.json"), encoding="utf-8") as handle:
            self.assertEqual(json.load(handle)["ausgabe"]["seiten"], 2)

    def test_placeholders_do_not_mask_the_page_limit_exit_code(self):
        source = SOURCE + "\n- [ERGÄNZEN: Aufgaben]\n"
        code, err = self.run_main_with_pages("letter", 2, source_text=source)
        self.assertEqual(code, 4)
        self.assertIn("1 offene Platzhalter", err)
        self.assertIn("Seitenlimit überschritten", err)

    def test_the_placeholder_marker_is_the_one_of_the_quality_gate_script(self):
        # the export and cv_checks.py (CV-01, BR-01) count the same marker; a new marker has to be added in both
        sys.path.insert(0, os.path.join(ROOT, "scripts"))
        sys.path.insert(0, os.path.dirname(EXPORTER))
        import cv_checks
        import export_document
        self.assertEqual(export_document.OPEN_PLACEHOLDER.pattern, cv_checks.PLACEHOLDER.pattern)

    def test_page_limits_differ_by_document_type(self):
        self.assertEqual(self.run_main_with_pages("cv", 2)[0], 0)
        self.assertEqual(self.run_main_with_pages("cv", 3)[0], 4)
        self.assertEqual(self.run_main_with_pages("letter", 1)[0], 0)
        studium = SOURCE.replace("format: chronologisch", "variante: studium")
        self.assertEqual(self.run_main_with_pages("motivation", 2, studium)[0], 0)
        self.assertEqual(self.run_main_with_pages("motivation", 3, studium)[0], 4)

    def test_job_variants_of_the_motivation_letter_are_limited_to_one_page(self):
        # MS-7: stelle (the default), quereinstieg and berufseinstieg are one page
        for variante in ("stelle", "quereinstieg", "berufseinstieg", None):
            text = SOURCE.replace("format: chronologisch", f"variante: {variante}" if variante else "format: chronologisch")
            self.assertEqual(self.run_main_with_pages("motivation", 1, text)[0], 0, variante)
            self.assertEqual(self.run_main_with_pages("motivation", 2, text)[0], 4, variante)

    def test_einseitig_is_recognised_in_any_spelling(self):
        # a stray capital or space in the frontmatter must not lose the one-page limit
        for value in ("Einseitig", '"einseitig "'):
            text = SOURCE.replace("format: chronologisch", f"format: {value}")
            self.assertEqual(self.run_main_with_pages("cv", 2, text)[0], 4, value)

    def test_overrun_is_in_the_manifest_even_with_a_content_difference(self):
        # R-10.10: code 2 wins, but the breach stays visible to the quality gate
        from unittest import mock
        sys.path.insert(0, os.path.dirname(EXPORTER))
        import export_document
        with mock.patch.object(export_document, "compare_content", lambda body, text: ({"fehlt": 1}, {})):
            code, _ = self.run_main_with_pages("letter", 2)
        self.assertEqual(code, 2)
        with open(os.path.join(self.out, "Test_Doc.pdf.manifest.json"), encoding="utf-8") as handle:
            ausgabe = json.load(handle)["ausgabe"]
        self.assertTrue(ausgabe["seitenlimit_ueberschritten"])
        self.assertEqual(ausgabe["seitenlimit"], 1)

    def test_einseitig_cv_is_limited_to_one_page(self):
        # LB-2.3: the format einseitig is one page
        einseitig = SOURCE.replace("format: chronologisch", "format: einseitig")
        self.assertEqual(self.run_main_with_pages("cv", 1, einseitig)[0], 0)
        self.assertEqual(self.run_main_with_pages("cv", 2, einseitig)[0], 4)

    def test_out_dir_inside_the_plugin_repository_is_refused(self):
        # LS-5.7, LS-2.4: nothing is written into the plugin repository
        result = subprocess.run(
            [sys.executable, EXPORTER, self.source, "--format", "md", "--doc-type", "cv", "--out-dir", os.path.join(ROOT, "export-test"), "--name", "X"],
            capture_output=True,
            text=True,
        )
        self.assertEqual(result.returncode, 1)
        self.assertIn("Plugin-Repository", result.stderr)
        self.assertFalse(os.path.exists(os.path.join(ROOT, "export-test")))

    def test_missing_grundlage_hash_is_reported_on_stdout(self):
        # R-10.3: the report names the missing hash
        with open(self.source, "w", encoding="utf-8") as handle:
            handle.write(SOURCE.replace("grundlage_hash: " + "0" * 64 + "\n", ""))
        result = self.export("md")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("grundlage_hash: fehlt", result.stdout)

    def test_manifest_records_environment_variables(self):
        # LV-6.3, LV-7.1: PLUGIN_REVISION and EXPORT_IMAGE_ID land in the manifest
        env = dict(os.environ, PLUGIN_REVISION="abc123", EXPORT_IMAGE_ID="sha256:img")
        subprocess.run(
            [sys.executable, EXPORTER, self.source, "--format", "md", "--doc-type", "cv", "--out-dir", self.out, "--name", "Test_Lebenslauf"],
            capture_output=True, text=True, env=env,
        )
        with open(os.path.join(self.out, "Test_Lebenslauf.md.manifest.json"), encoding="utf-8") as handle:
            umgebung = json.load(handle)["umgebung"]
        self.assertEqual(umgebung, {"plugin_revision": "abc123", "image_id": "sha256:img"})

    def test_tex_export_writes_source_without_lualatex(self):
        # tex is the LaTeX source of the PDF; it needs pandoc only and passes the content check
        result = self.export("tex")
        self.assertEqual(result.returncode, 0, result.stderr)
        with open(os.path.join(self.out, "Test_Lebenslauf.tex"), encoding="utf-8") as handle:
            tex = handle.read()
        self.assertIn("\\begin{document}", tex)
        self.assertIn("Latenz von 800 ms", tex)
        # the layout filter maps the entry shapes of the Lebenslauf spec (LB-4.1, LB-5.1) to the template macros
        flat = " ".join(tex.split())  # pandoc wraps long lines
        self.assertIn("\\cvheader{Lebenslauf}{Max Mustermann}", flat)
        self.assertIn("\\cvstation{04/2021", flat)

    def test_letter_gets_its_own_page_geometry_and_no_cv_filter(self):
        # the Anschreiben uses 25 mm left (AS-3.3) and no Lebenslauf macros
        result = subprocess.run(
            [sys.executable, EXPORTER, self.source, "--format", "tex", "--doc-type", "letter", "--out-dir", self.out, "--name", "Test_Brief"],
            capture_output=True,
            text=True,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        with open(os.path.join(self.out, "Test_Brief.tex"), encoding="utf-8") as handle:
            tex = handle.read()
        self.assertIn("left=25mm", tex)
        self.assertNotIn("\\cvstation{", tex)

    def test_readback_macros_cover_the_three_layout_macros(self):
        sys.path.insert(0, os.path.dirname(EXPORTER))
        import export_document

        self.assertEqual(export_document.READBACK_MACROS.count("newcommand"), 3)

    def test_gap_entry_and_short_row_render_in_the_grid(self):
        # LB-5.1: a two-part station ("Zeitraum | Elternzeit") is a station with an empty employer, not a
        # fallback heading; LB-6.1/LB-11.9: a two-word row keeps the no-break space out of the label split
        with open(self.source, "w", encoding="utf-8") as handle:
            handle.write(SOURCE.replace(
                "## Sprachkenntnisse (GER-Niveau)\n",
                "### 01/2019 – 03/2021 | Elternzeit\n\n## Ausbildung und Studium\n- 2018: Informatik\n\n## Sprachkenntnisse (GER-Niveau)\n",
            ))
        result = self.export("tex")
        self.assertEqual(result.returncode, 0, result.stderr)
        with open(os.path.join(self.out, "Test_Lebenslauf.tex"), encoding="utf-8") as handle:
            flat = " ".join(handle.read().split())
        self.assertIn("\\cvstation{01/2019 -- 03/2021}{Elternzeit}{}", flat)
        self.assertNotIn("\\textbar", flat)
        self.assertIn("\\cvrow{2018}{Informatik}", flat)
        self.assertIn("\\textbf{Deutsch}: Muttersprache", flat)

    @unittest.skipUnless(python_docx_available(), "python-docx is not installed")
    def test_docx_export_keeps_headings_in_order_and_bullets(self):
        # Scenario I/J: standard headings in source order, no bullet lost, no table
        import docx

        result = self.export("docx")
        self.assertEqual(result.returncode, 0, result.stderr)
        document = docx.Document(os.path.join(self.out, "Test_Lebenslauf.docx"))
        texts = [p.text for p in document.paragraphs]
        positions = [next(i for i, t in enumerate(texts) if t.startswith(h)) for h in HEADINGS]
        self.assertEqual(positions, sorted(positions))
        self.assertEqual(len(document.tables), 0)
        self.assertTrue(any("Einführung von Code-Reviews" in t for t in texts))

    @unittest.skipUnless(have("lualatex"), "LuaLaTeX is not installed: PDF export and page check not run")
    def test_pdf_export_has_extractable_text(self):
        # CG-12.4
        result = self.export("pdf")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue(have("pdftotext"))
        text = subprocess.run(["pdftotext", os.path.join(self.out, "Test_Lebenslauf.pdf"), "-"], capture_output=True, text=True).stdout
        self.assertIn("Berufserfahrung", text)

    @unittest.skipUnless(have("lualatex"), "LuaLaTeX is not installed: PDF export and page check not run")
    def test_pdf_text_keeps_figures_symbols_and_the_reading_order_over_several_pages(self):
        # F-04, groups B, G: the text a parser extracts has the station line in order, no figure or symbol is lost, and the
        # new check finds a swapped figure in the same PDF (not only in a simulated text)
        stations = "".join(
            f"\n### {2000 + i:04d} – {2001 + i:04d} | Position {i} | Beispiel AG {i}, Hamburg\n"
            + "".join(f"- Aufgabe {j}: Plattform mit 40 % weniger Störungen betrieben; C++ und C# eingesetzt, siehe [Seite](https://beispiel.test/a-{i}).\n" for j in range(1, 5))
            for i in range(1, 7))
        text = SOURCE.replace("## Berufserfahrung\n", "## Berufserfahrung\n" + stations, 1)
        with open(self.source, "w", encoding="utf-8") as handle:
            handle.write(text)
        result = self.export("pdf")
        self.assertIn(result.returncode, (0, 4), result.stderr)  # 4: more than two pages, which is not the subject here
        self.assertIn("Inhaltsprüfung: ok", result.stdout)
        pdf = os.path.join(self.out, "Test_Lebenslauf.pdf")
        sys.path.insert(0, os.path.dirname(EXPORTER))
        import export_document

        extracted = export_document.strip_page_numbers(subprocess.run(["pdftotext", pdf, "-"], capture_output=True, text=True).stdout)
        flat = " ".join(extracted.split())
        self.assertRegex(flat, r"2003 – 2004\s+Position 3\s+Beispiel AG 3, Hamburg")
        self.assertEqual(export_document.fact_differences(export_document.read_document(self.source)[1], extracted), [])
        self.assertTrue(export_document.fact_differences(export_document.read_document(self.source)[1].replace("40 % weniger", "weniger 40 %", 1), extracted))
        with open(os.path.join(self.out, "Test_Lebenslauf.pdf.manifest.json"), encoding="utf-8") as handle:
            self.assertEqual(json.load(handle)["pruefung"]["zeichen_und_stichpunkte"], "ok")

    def freeze(self):
        with open(self.source, "w", encoding="utf-8") as handle:
            handle.write(SOURCE.replace("format: chronologisch\n", "format: chronologisch\nauspraegung: basis\nversion: 1\nversion_datum: 2026-01-01\n"))

    def test_export_writes_a_manifest_that_pins_source_and_output(self):
        # lebenslauf-versionen LV-6: the manifest names the source hash, the epoch and the output hash
        self.freeze()
        self.assertEqual(self.export("tex").returncode, 0)
        with open(os.path.join(self.out, "Test_Lebenslauf.tex.manifest.json"), encoding="utf-8") as handle:
            manifest = json.load(handle)
        self.assertEqual(manifest["frontmatter"]["version"], "1")
        self.assertEqual(manifest["source_date_epoch"], 1767225600)
        self.assertTrue(manifest["byte_reproduzierbar"])
        self.assertEqual(len(manifest["quelle"]["sha256"]), 64)
        self.assertEqual(len(manifest["ausgabe"]["sha256"]), 64)

    def test_invalid_version_datum_is_an_error_not_a_traceback(self):
        with open(self.source, "w", encoding="utf-8") as handle:
            handle.write(SOURCE.replace("format: chronologisch\n", "format: chronologisch\nversion_datum: 2026-02-30\n"))
        result = self.export("tex")
        self.assertEqual(result.returncode, 1)
        self.assertIn("version_datum", result.stderr)
        self.assertNotIn("Traceback", result.stderr)

    def test_invalid_version_datum_only_warns_for_formats_without_a_date(self):
        # DOCX, HTML and MD carry no PDF date: like a missing date, a bad one is reported and the export goes on
        with open(self.source, "w", encoding="utf-8") as handle:
            handle.write(SOURCE.replace("format: chronologisch\n", "format: chronologisch\nversion_datum: '2026-02-30'\n"))
        for fmt in ("html", "md"):
            result = self.export(fmt)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("version_datum", result.stderr)
            self.assertNotIn("Traceback", result.stderr)

    def test_quoted_version_datum_is_read_like_the_bare_date(self):
        # YAML quotes around the date are not part of the value
        for quoted in ('"2026-01-01"', "'2026-01-01'"):
            with open(self.source, "w", encoding="utf-8") as handle:
                handle.write(SOURCE.replace("format: chronologisch\n", f"format: chronologisch\nversion_datum: {quoted}\n"))
            for fmt in ("html", "md", "tex"):
                result = self.export(fmt)
                self.assertEqual(result.returncode, 0, f"{quoted} {fmt}: {result.stderr}")
            with open(os.path.join(self.out, "Test_Lebenslauf.tex.manifest.json"), encoding="utf-8") as handle:
                manifest = json.load(handle)
            self.assertEqual(manifest["source_date_epoch"], 1767225600)
            self.assertEqual(manifest["frontmatter"]["version_datum"], "2026-01-01")

    def test_reproduce_confirms_an_unchanged_export_and_rejects_a_changed_source(self):
        self.freeze()
        self.export("tex")
        manifest = os.path.join(self.out, "Test_Lebenslauf.tex.manifest.json")
        ok = subprocess.run([sys.executable, EXPORTER, "--reproduce", manifest], capture_output=True, text=True)
        self.assertEqual(ok.returncode, 0, ok.stderr)
        self.assertIn("byte-identisch", ok.stdout)
        with open(self.source, "a", encoding="utf-8") as handle:
            handle.write("\nNachtrag\n")
        changed = subprocess.run([sys.executable, EXPORTER, "--reproduce", manifest], capture_output=True, text=True)
        self.assertEqual(changed.returncode, 3)
        self.assertIn("Quelle verändert", changed.stderr)

    def test_reproduce_accepts_a_manifest_with_a_non_ascii_name(self):
        # a manifest written before 0.2.0 may carry an umlaut in `name` and `ausgabe.datei`; the rebuilt
        # file is found under today's transliteration (LV-6)
        self.freeze()
        self.export("tex")
        manifest = os.path.join(self.out, "Test_Lebenslauf.tex.manifest.json")
        with open(manifest, encoding="utf-8") as handle:
            data = json.load(handle)
        data["name"] = "Test_Lébenslauf"  # transliterates to the same file name as the export above
        data["ausgabe"]["datei"] = "Test_Lébenslauf.tex"
        with open(manifest, "w", encoding="utf-8") as handle:
            json.dump(data, handle, ensure_ascii=False)
        result = subprocess.run([sys.executable, EXPORTER, "--reproduce", manifest], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("byte-identisch", result.stdout)

    def test_frozen_version_exports_next_to_the_arbeitsstand_by_default(self):
        # lokale-struktur LS-3, lebenslauf-versionen: versionen/vNNN_<datum>.md -> ../export/vNNN_<datum>.<ext>
        folder = os.path.join(self.tmp, "lebenslauf-basis")
        os.makedirs(os.path.join(folder, "versionen"))
        frozen = os.path.join(folder, "versionen", "v001_2026-01-01.md")
        with open(frozen, "w", encoding="utf-8") as handle:
            handle.write(SOURCE.replace("format: chronologisch\n", "format: chronologisch\nauspraegung: lebenslauf-basis\nversion: 1\nversion_datum: 2026-01-01\n"))
        result = subprocess.run([sys.executable, EXPORTER, frozen, "--format", "md", "--doc-type", "cv"], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue(os.path.isfile(os.path.join(folder, "export", "v001_2026-01-01.md")), result.stdout)
        self.assertFalse(os.path.exists(os.path.join(folder, "versionen", "export")))

    def test_file_names_are_transliterated_to_ascii(self):
        # R-4: ä→ae, ö→oe, ü→ue, ß→ss, spaces to hyphens, hyphen and underscore kept
        sys.path.insert(0, os.path.dirname(EXPORTER))
        import export_document

        self.assertEqual(export_document.ascii_name("Jörg_Müller-Straße Lebenslauf"), "Joerg_Mueller-Strasse-Lebenslauf")
        # accented and non-German letters are transliterated, not dropped
        self.assertEqual(export_document.ascii_name("René_Dupont_Lebenslauf"), "Rene_Dupont_Lebenslauf")
        self.assertEqual(export_document.ascii_name("Zoë_Łukasz_Ørsted"), "Zoe_Lukasz_Orsted")
        self.assertEqual(export_document.ascii_name("François_Çelik"), "Francois_Celik")
        result = subprocess.run(
            [sys.executable, EXPORTER, self.source, "--format", "md", "--doc-type", "cv", "--out-dir", self.out, "--name", "Max_Müller_Lebenslauf"],
            capture_output=True, text=True,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue(os.path.isfile(os.path.join(self.out, "Max_Mueller_Lebenslauf.md")))

    def test_product_names_do_not_break_across_lines(self):
        # cv.lua: heuristic (inner capital or digit) plus the frontmatter list keep_together; a plain pair stays
        with open(self.source, "w", encoding="utf-8") as handle:
            handle.write(SOURCE.replace("format: chronologisch\n", "format: chronologisch\nkeep_together: [Docker Compose]\n").replace(
                "- Einführung von Code-Reviews im Team mit fünf Entwicklern.",
                "- Deployments mit Argo CD, GitLab CI und Docker Compose automatisiert; Python und Kafka eingesetzt",
            ))
        result = self.export("tex")
        self.assertEqual(result.returncode, 0, result.stderr)
        with open(os.path.join(self.out, "Test_Lebenslauf.tex"), encoding="utf-8") as handle:
            flat = " ".join(handle.read().split())
        for joined in ("Argo~CD", "GitLab~CI", "Docker~Compose"):
            self.assertIn(joined, flat)
        self.assertIn("Python und Kafka", flat)

    def test_manifest_does_not_carry_the_dropped_situation_key(self):
        with open(self.source, "w", encoding="utf-8") as handle:
            handle.write(SOURCE.replace("format: chronologisch\n", "format: chronologisch\nsituation: Standard\n"))
        self.assertEqual(self.export("md").returncode, 0)
        with open(os.path.join(self.out, "Test_Lebenslauf.md.manifest.json"), encoding="utf-8") as handle:
            self.assertNotIn("situation", json.load(handle)["frontmatter"])

    @unittest.skipUnless(python_docx_available(), "python-docx is not installed")
    def test_docx_metadata_names_no_author(self):
        # decision D2: no tool or account name in the DOCX author fields
        import docx

        self.assertEqual(self.export("docx").returncode, 0)
        properties = docx.Document(os.path.join(self.out, "Test_Lebenslauf.docx")).core_properties
        self.assertEqual((properties.author, properties.last_modified_by), ("", ""))
        self.assertEqual(properties.title, "Test_Lebenslauf")

    @unittest.skipUnless(have("lualatex"), "LuaLaTeX is not installed: PDF reproducibility not run")
    def test_frozen_pdf_is_byte_identical_on_a_second_export(self):
        self.freeze()
        self.export("pdf")
        first = open(os.path.join(self.out, "Test_Lebenslauf.pdf"), "rb").read()
        self.export("pdf")
        self.assertEqual(first, open(os.path.join(self.out, "Test_Lebenslauf.pdf"), "rb").read())


class OrderTests(unittest.TestCase):
    """EXP-002: the export compares the order of the words, not only the multiset (CG-12.2)."""

    def setUp(self):
        sys.path.insert(0, os.path.dirname(EXPORTER))
        import export_document

        self.mod = export_document

    def test_equal_text_has_no_difference_even_with_a_line_broken_hyphenated_word(self):
        self.assertIsNone(self.mod.first_order_difference("## Team\nDie E-Mail an das Team", "Team\nDie E-\nMail an das Team"))

    def test_two_swapped_sections_name_the_first_heading_out_of_place(self):
        source = "## Berufserfahrung\nPython\n## Ausbildung\nBachelor"
        exported = "Ausbildung\nBachelor\nBerufserfahrung\nPython"
        self.assertEqual(self.mod.compare_content(source, exported), (self.mod.Counter(), self.mod.Counter()))
        self.assertEqual(self.mod.first_order_difference(source, exported), "Ausbildung")

    def test_moved_header_fields_and_station_headings_are_not_a_difference(self):
        # the CV layout puts the name before "Persönliche Daten" and lays out the station lines of level 3 as a grid
        source = "## Persönliche Daten\nMax Mustermann\n## Berufserfahrung\n### 04/2021 | Entwickler | Firma\n- Python"
        exported = "Max Mustermann\nPersönliche Daten\nBerufserfahrung\nFirma Entwickler 04/2021\nPython"
        self.assertIsNone(self.mod.first_order_difference(source, exported))

    def test_a_heading_word_in_the_running_text_can_hide_a_swap_a_known_limit(self):
        source = "## Berufserfahrung\nIch habe eine Ausbildung gemacht\n## Ausbildung\nBachelor"
        swapped = "Ausbildung\nBachelor\nBerufserfahrung\nIch habe eine Ausbildung gemacht"
        self.assertIsNone(self.mod.first_order_difference(source, swapped))

    def test_the_export_warns_about_the_order_without_changing_the_exit_code(self):
        tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmp, ignore_errors=True)
        source = os.path.join(tmp, "lebenslauf.md")
        with open(source, "w", encoding="utf-8") as handle:
            handle.write(SOURCE)
        reordered = SOURCE.split("---\n", 2)[2].split("\n## ")
        reordered = "\n## ".join([reordered[0], reordered[2], reordered[1], *reordered[3:]])
        import io
        from contextlib import redirect_stderr, redirect_stdout
        from unittest import mock

        out, err = io.StringIO(), io.StringIO()
        with mock.patch.object(self.mod, "require", return_value="pandoc"), mock.patch.object(self.mod, "export_html", return_value=(reordered, None)), mock.patch.object(self.mod, "structure_notes", return_value=[]), redirect_stdout(out), redirect_stderr(err):
            code = self.mod.main([source, "--format", "html", "--doc-type", "cv", "--out-dir", os.path.join(tmp, "out"), "--no-manifest"])
        self.assertEqual(code, 0, err.getvalue())
        self.assertIn("Inhaltsprüfung: ok", out.getvalue())
        self.assertIn("Reihenfolge: Abweichung", out.getvalue())
        self.assertIn("Reihenfolge der Überschriften weicht ab", err.getvalue())

    @unittest.skipUnless(have("pandoc"), "pandoc is not installed")
    def test_a_normal_export_prints_no_order_line(self):
        tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmp, ignore_errors=True)
        source = os.path.join(tmp, "lebenslauf.md")
        with open(source, "w", encoding="utf-8") as handle:
            handle.write(SOURCE)
        for fmt in ("html", "docx", "tex"):
            result = subprocess.run([sys.executable, EXPORTER, source, "--format", fmt, "--doc-type", "cv", "--out-dir", os.path.join(tmp, "out")], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, f"{fmt}: {result.stderr}")
            self.assertNotIn("Reihenfolge", result.stdout + result.stderr, fmt)


BULLETS = """\
## Berufserfahrung

### 04/2021 – heute | Senior Backend-Entwickler | Musterfirma GmbH, Berlin
- Latenz der Pipeline von 800 ms auf 120 ms gesenkt.
- Umsatz um 40 % gesteigert; C++ und C# eingesetzt.
- Code-Reviews im Team mit fünf Entwicklern eingeführt.
"""


class FactDifferenceTests(unittest.TestCase):
    """F-04: what the word multiset cannot see. Groups A (content kept), F (no changed fact), G (reading order)."""

    def setUp(self):
        sys.path.insert(0, os.path.dirname(EXPORTER))
        import export_document

        self.mod = export_document

    def differences(self, exported, source=BULLETS):
        return self.mod.fact_differences(source, exported)

    def test_the_old_check_misses_swapped_figures_lost_symbols_and_swapped_bullets(self):
        # reproduction of F-04: each of these is equal to the source as a multiset of words
        for damaged in (BULLETS.replace("von 800 ms auf 120 ms", "von 120 ms auf 800 ms"),
                        BULLETS.replace("40 %", "40").replace("C++", "C").replace("C#", "C"),
                        "\n".join(BULLETS.splitlines()[:3] + BULLETS.splitlines()[:2:-1]) + "\n"):
            self.assertEqual(self.mod.compare_content(BULLETS, damaged), (self.mod.Counter(), self.mod.Counter()))

    def test_swapped_figures_in_a_bullet_are_found(self):
        found = self.differences(BULLETS.replace("von 800 ms auf 120 ms", "von 120 ms auf 800 ms"))
        self.assertEqual(len(found), 1)
        self.assertIn("latenz der pipeline", found[0])

    def test_two_swapped_bullets_are_found(self):
        lines = BULLETS.splitlines()
        swapped = "\n".join(lines[:3] + [lines[4], lines[3]] + lines[5:])
        self.assertTrue(self.differences(swapped))

    def test_a_lost_percent_sign_and_plus_signs_are_found(self):
        found = self.differences(BULLETS.replace("40 %", "40").replace("C++", "C"))
        self.assertIn("%", found[0])
        self.assertIn("+", found[0])

    def test_an_added_symbol_is_found(self):
        self.assertIn("hinzu", self.differences(BULLETS + "Zusatz 5 %\n")[0])

    def test_layout_changes_are_not_a_difference(self):
        # line and page breaks, a hyphen at the line end, a bullet sign, link targets and bold marks (valid-looking layout)
        wrapped = ("Berufserfahrung\n04/2021 – heute Senior Backend-Entwickler Musterfirma GmbH, Berlin\n"
                   "• Latenz der Pipeline von 800 ms\nauf 120 ms gesenkt.\n\f• Umsatz um 40 % gesteigert; C++ und C# eingesetzt.\n"
                   "• Code-\nReviews im Team mit fünf Entwicklern eingeführt.\n")
        self.assertEqual(self.differences(wrapped), [])
        linked = "- Details unter [Beispiel](https://beispiel.test/a-b) **veröffentlicht**.\n"
        self.assertEqual(self.differences("Details unter Beispiel veröffentlicht.", linked), [])

    def test_other_words_in_other_places_of_the_page_do_not_matter_for_the_items(self):
        # header fields and the station lines of level 3 are laid out freely; only the items are held to their order
        moved = "Musterfirma GmbH, Berlin\nSenior Backend-Entwickler 04/2021 – heute\n" + "\n".join(BULLETS.splitlines()[3:]) + "\nBerufserfahrung\n"
        self.assertEqual(self.differences(moved), [])

    def test_the_source_without_list_items_or_symbols_has_nothing_to_find(self):
        self.assertEqual(self.mod.fact_differences("## Titel\nEin Absatz ohne Liste.\n", "Ein Absatz ohne Liste. Titel"), [])


class StructureAndManifestTests(unittest.TestCase):
    """F-04: structure notes of DOCX and HTML, the manifest key `pruefung`, the pandoc option. Groups C, G, H, I, J."""

    TABLE_SOURCE = (
        "---\ngrundlage_hash: 0000000000000000000000000000000000000000000000000000000000000000\n---\n\n# Lebenslauf\n\n"
        "## Berufserfahrung\n\n| Zeitraum | Position |\n|---|---|\n| 04/2021 – heute | Senior Backend-Entwickler |\n\n"
        "- Latenz von 800 ms auf 120 ms gesenkt.\n")

    def setUp(self):
        sys.path.insert(0, os.path.dirname(EXPORTER))
        import export_document

        self.mod = export_document
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    def reading(self, text):
        """export_html that writes the real file but reads back `text`: the export came out damaged, the file exists."""
        real = self.mod.export_html
        return lambda *args: (real(*args)[0] and text, None)

    def run_export(self, text, fmt, name="lebenslauf"):
        import io
        from contextlib import redirect_stderr, redirect_stdout

        source = os.path.join(self.tmp, name + ".md")
        with open(source, "w", encoding="utf-8") as handle:
            handle.write(text)
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = self.mod.main([source, "--format", fmt, "--doc-type", "cv", "--out-dir", os.path.join(self.tmp, "out")])
        manifest = os.path.join(self.tmp, "out", f"{name}.{fmt}.manifest.json")
        data = json.load(open(manifest, encoding="utf-8")) if os.path.exists(manifest) else None
        return code, out.getvalue(), err.getvalue(), data

    @unittest.skipUnless(have("pandoc"), "pandoc is not installed")
    @unittest.skipUnless(python_docx_available(), "python-docx is not installed")
    def test_a_table_in_the_docx_is_read_as_content_and_reported_as_structure(self):
        # before F-04 the cells of the table counted as missing words (exit 2) and the table itself was not named
        code, out, err, manifest = self.run_export(self.TABLE_SOURCE, "docx")
        self.assertEqual(code, 0, err)
        self.assertIn("Inhaltsprüfung: ok", out)
        self.assertIn("Struktur: 1 Tabelle(n)", err)
        self.assertEqual(manifest["pruefung"]["inhalt"], "ok")
        self.assertTrue(any("Tabelle" in note for note in manifest["pruefung"]["struktur_hinweise"]))

    @unittest.skipUnless(have("pandoc"), "pandoc is not installed")
    def test_html_reports_a_table_and_keeps_the_exit_code(self):
        code, out, err, manifest = self.run_export(self.TABLE_SOURCE, "html")
        self.assertEqual(code, 0, err)
        self.assertTrue(any("Tabelle" in note for note in manifest["pruefung"]["struktur_hinweise"]))

    @unittest.skipUnless(have("pandoc"), "pandoc is not installed")
    @unittest.skipUnless(python_docx_available(), "python-docx is not installed")
    def test_a_normal_docx_and_html_have_no_structure_note(self):
        for fmt in ("docx", "html"):
            code, out, err, manifest = self.run_export(SOURCE, fmt, name="normal-" + fmt)
            self.assertEqual(code, 0, err)
            self.assertEqual(manifest["pruefung"]["struktur_hinweise"], [], fmt)
            self.assertNotIn("Struktur:", out + err)

    def test_headings_that_do_not_match_the_source_are_named(self):
        # the same notes from a DOCX that lost a heading level: simulated with an HTML file
        from pathlib import Path

        html_file = Path(self.tmp) / "x.html"
        html_file.write_text("<h1>Lebenslauf</h1><h3>Persönliche Daten</h3>", encoding="utf-8")
        notes = self.mod.structure_notes("html", html_file, "# Lebenslauf\n\n## Persönliche Daten\n")
        self.assertIn("Überschriften weichen von der Quelle ab", notes[0])
        self.assertIn("H2×1", notes[0])

    @unittest.skipUnless(have("pandoc"), "pandoc is not installed")
    def test_the_manifest_records_the_checks_of_the_run_and_what_they_do_not_cover(self):
        code, out, err, manifest = self.run_export(SOURCE, "html")
        checks = manifest["pruefung"]
        self.assertEqual(checks["gegen"], "quelle")
        self.assertEqual((checks["inhalt"], checks["zeichen_und_stichpunkte"], checks["reihenfolge_ueberschriften"]), ("ok", "ok", "ok"))
        self.assertEqual((checks["platzhalter_offen"], checks["grundlage_hash_vorhanden"]), (0, True))
        self.assertNotIn("ats", json.dumps(manifest).lower())

    @unittest.skipUnless(have("pandoc"), "pandoc is not installed")
    def test_markdown_is_a_copy_and_is_recorded_as_not_checked(self):
        code, out, err, manifest = self.run_export(SOURCE, "md")
        self.assertEqual(code, 0, err)
        self.assertEqual(manifest["pruefung"]["inhalt"], "nicht_geprueft")

    @unittest.skipUnless(have("pandoc"), "pandoc is not installed")
    def test_a_failed_check_is_in_the_manifest_and_a_damaged_figure_gives_exit_2(self):
        from unittest import mock

        damaged = SOURCE.split("---\n", 2)[2].replace("von 800 ms auf 120 ms", "von 120 ms auf 800 ms")
        with mock.patch.object(self.mod, "export_html", self.reading(damaged)):
            code, out, err, manifest = self.run_export(SOURCE, "html")
        self.assertEqual(code, 2, err)
        self.assertIn("Inhaltsabweichung: Stichpunkt im Export", err)
        self.assertEqual(manifest["pruefung"]["inhalt"], "abweichung")
        self.assertEqual(manifest["pruefung"]["zeichen_und_stichpunkte"], "abweichung")
        self.assertEqual(manifest["pruefung"]["reihenfolge_ueberschriften"], "nicht_geprueft")

    @unittest.skipUnless(have("pandoc"), "pandoc is not installed")
    def test_an_added_statement_is_found_as_an_extra_word(self):
        from unittest import mock

        added = SOURCE.split("---\n", 2)[2] + "\n- Teamleitung für zwölf Mitarbeiter übernommen.\n"
        with mock.patch.object(self.mod, "export_html", self.reading(added)):
            code, out, err, manifest = self.run_export(SOURCE, "html")
        self.assertEqual(code, 2)
        self.assertIn("Zusätzlich im Export", err)

    @unittest.skipUnless(have("pandoc"), "pandoc is not installed")
    def test_reproduce_still_confirms_a_docx_and_an_html_export(self):
        for fmt in ("html", "docx") if python_docx_available() else ("html",):
            code, out, err, manifest = self.run_export(SOURCE, fmt, name="rep-" + fmt)
            self.assertEqual(code, 0, err)
            self.assertEqual(self.mod.main(["--reproduce", os.path.join(self.tmp, "out", f"rep-{fmt}.{fmt}.manifest.json")]), 0, fmt)

    @unittest.skipUnless(have("pandoc"), "pandoc is not installed")
    @unittest.skipUnless(python_docx_available(), "python-docx is not installed")
    def test_a_link_is_content_by_its_text_and_not_by_its_target(self):
        # before F-04 the words of the URL counted as missing in every format and the content check failed
        source = SOURCE.replace("Python und Kafka;", "Python und Kafka (siehe [Beispielseite](https://beispiel.test/a-b));", 1)
        for fmt in ("docx", "html"):
            code, out, err, manifest = self.run_export(source, fmt, name="link-" + fmt)
            self.assertEqual(code, 0, fmt + err)
            self.assertEqual(manifest["pruefung"]["inhalt"], "ok", fmt)

    def test_the_pandoc_option_follows_the_version(self):
        self.assertEqual(self.mod.embed_option("pandoc 2.17.1.1"), "--self-contained")
        self.assertEqual(self.mod.embed_option("pandoc 2.19.2"), "--embed-resources")
        self.assertEqual(self.mod.embed_option("pandoc 3.1.3"), "--embed-resources")
        self.assertEqual(self.mod.embed_option(None), "--embed-resources")

    def test_missing_export_data_and_a_bad_format_are_errors_not_tracebacks(self):
        import io
        from contextlib import redirect_stderr, redirect_stdout

        err = io.StringIO()
        with redirect_stdout(io.StringIO()), redirect_stderr(err):
            code = self.mod.main([os.path.join(self.tmp, "missing.md"), "--format", "md", "--doc-type", "cv", "--out-dir", self.tmp + "/o"])
        self.assertEqual(code, 1)
        self.assertIn("Fehler", err.getvalue())


class PageNumberTests(unittest.TestCase):
    def strip(self, text):
        sys.path.insert(0, os.path.dirname(EXPORTER))
        import export_document

        return export_document.strip_page_numbers(text)

    def test_year_in_the_period_column_is_not_taken_for_a_page_number(self):
        # page 2 ends with a period line "2019": content, not the page number
        self.assertEqual(self.strip("Berufserfahrung\n1\n\fAusbildung\n2019\n"), "Berufserfahrung\nAusbildung\n2019")

    def test_page_number_at_the_end_of_its_page_is_dropped(self):
        self.assertEqual(self.strip("Berufserfahrung\n1\n\fAusbildung\n2\n\n"), "Berufserfahrung\nAusbildung")

    def test_single_page_without_a_number_is_unchanged(self):
        self.assertEqual(self.strip("Berufserfahrung\nAusbildung"), "Berufserfahrung\nAusbildung")


class CompareContentTests(unittest.TestCase):
    def test_hyphen_at_line_break_is_not_a_difference(self):
        # A PDF breaks "E-Mail" or "Kubernetes-Cluster" at the hyphen; the words are the same
        sys.path.insert(0, os.path.dirname(EXPORTER))
        import export_document

        missing, extra = export_document.compare_content("E-Mail und Kubernetes-Cluster", "E-\nMail und Kubernetes-\nCluster")
        self.assertEqual((dict(missing), dict(extra)), ({}, {}))

    def test_missing_word_is_still_reported(self):
        sys.path.insert(0, os.path.dirname(EXPORTER))
        import export_document

        missing, _ = export_document.compare_content("Python und Kafka", "Python")
        self.assertEqual(sorted(missing), ["kafka", "und"])


if __name__ == "__main__":
    unittest.main()

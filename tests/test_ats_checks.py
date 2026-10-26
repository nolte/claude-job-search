"""Tests of the ATS checks `scripts/ats_checks.py` and their place in `qualitaet.py` (spec qualitaetssicherung QS-13).

The checks run on synthetic exports: PDFs built byte by byte (so no LaTeX is needed), DOCX files built with python-docx or
by changing one part of a real pandoc export, and the real exporter output for HTML and DOCX. What is tested is what the
checks find in a planted defect and that a clean export passes, not what any tracking system would do. All data is fictional.
"""

import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import unittest.mock
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "scripts"))

import ats_checks as ats  # noqa: E402
import qualitaet as qa  # noqa: E402
import test_qualitaet as tq  # noqa: E402

POPPLER = all(shutil.which(t) for t in ("pdftotext", "pdfinfo", "pdffonts", "pdfimages"))
PANDOC = shutil.which("pandoc") is not None
try:
    import docx
except ImportError:  # pragma: no cover
    docx = None

CV = """# Lebenslauf

## Persönliche Daten
Maria Beispiel, Berlin, maria.beispiel@beispiel-firma.example, +49 30 1234567

## Berufserfahrung

### 03/2022 – heute | Teamleiterin Plattform | Musterfirma GmbH, Berlin
- Kubernetes Plattform betrieben und Bereitstellungszeit deutlich gesenkt
- Fachliche Leitung des Plattform Teams übernommen

### 09/2017 – 02/2022 | Softwareentwicklerin | Beispiel AG, Hamburg
- REST Schnittstellen in Python entwickelt und Partnersysteme angebunden

## Ausbildung und Studium
- 10/2009 – 09/2013: Bachelor Informatik, Universität Beispiel, Hamburg
"""

# What pdftotext reads back from the page of a clean export: heading lines, station lines, bullets.
CV_LINES = [
    (72, 780, "Lebenslauf"), (72, 760, "Persönliche Daten"),
    (72, 745, "Maria Beispiel, Berlin, maria.beispiel@beispiel-firma.example, +49 30 1234567"),
    (72, 720, "Berufserfahrung"),
    (72, 700, "03/2022 – heute"), (72, 686, "Teamleiterin Plattform"), (72, 672, "Musterfirma GmbH, Berlin"),
    (72, 658, "Kubernetes Plattform betrieben und Bereitstellungszeit deutlich gesenkt"),
    (72, 644, "Fachliche Leitung des Plattform Teams übernommen"),
    (72, 620, "09/2017 – 02/2022"), (72, 606, "Softwareentwicklerin"), (72, 592, "Beispiel AG, Hamburg"),
    (72, 578, "REST Schnittstellen in Python entwickelt und Partnersysteme angebunden"),
    (72, 550, "Ausbildung und Studium"),
    (72, 536, "10/2009 – 09/2013: Bachelor Informatik, Universität Beispiel, Hamburg"),
]


def positioned_pdf(pages, image=False, font="/Type /Font /Subtype /Type1 /BaseFont /Helvetica", width=595):
    """A valid PDF whose text stands at given places. `pages` is a list of pages, each a list of (x, y, text)."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", None, f"<< {font} /Encoding /WinAnsiEncoding >>"]
    if image:
        objects.append("<< /Type /XObject /Subtype /Image /Width 1 /Height 1 /ColorSpace /DeviceGray /BitsPerComponent 8 /Length 1 >>\nstream\n\x80\nendstream")
    kids = []
    for runs in pages:
        number = len(objects) + 1
        kids.append(f"{number} 0 R")
        stream = "".join(f"BT /F1 11 Tf {x} {y} Td ({t.replace(chr(92), chr(92) * 2).replace('(', chr(92) + '(').replace(')', chr(92) + ')')}) Tj ET\n" for x, y, t in runs)
        if image:
            stream += "q 40 0 0 40 400 700 cm /Im1 Do Q\n"
        xobject = "/XObject << /Im1 4 0 R >> " if image else ""
        encoded = stream.encode("cp1252")
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {width} 842] /Resources << /Font << /F1 3 0 R >> {xobject}>> /Contents {number + 1} 0 R >>")
        objects.append(("STREAM", encoded))
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>"
    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        if isinstance(body, tuple):
            out += f"{number} 0 obj\n<< /Length {len(body[1])} >>\nstream\n".encode("latin-1") + body[1] + b"endstream\nendobj\n"
        else:
            out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    start = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode() + "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n".encode()
    return out


def write(tmp, name, content):
    path = os.path.join(tmp, name)
    with open(path, "wb") as handle:
        handle.write(content)
    return path


def messages(findings, check_id=None):
    return [(c, g, m) for c, g, m in findings if check_id in (None, c)]


STRUCTURE = {"tabellen": 0, "bilder": 0}


class TextTests(unittest.TestCase):
    def test_a_normal_text_has_no_finding(self):
        self.assertEqual(ats.text_quality("Kubernetes Plattform betrieben " * 12, "pdf"), [])

    def test_too_little_text_replacement_characters_ligatures_and_letter_spacing_are_found(self):
        self.assertEqual([g for _, g, _ in ats.text_quality("zwei Wörter", "pdf")], ["Blocker"])
        self.assertEqual([g for _, g, _ in ats.text_quality("zwei Wörter", "html")], [])  # html is no image of a page
        for defekt in ("Betrieb � der Plattform " * 10, "Betrieb (cid:12) der Plattform " * 10):
            self.assertIn("Blocker", [g for _, g, _ in ats.text_quality(defekt, "docx")])
        self.assertIn("Warnung", [g for _, g, _ in ats.text_quality("Zertiﬁzierung Kubernetes " * 20, "pdf")])
        self.assertIn("Warnung", [g for _, g, _ in ats.text_quality("Erfahrung k u b e r n e t e s Betrieb " * 10, "pdf")])
        self.assertEqual([m for _, _, m in ats.text_quality("Kenntnisse C R M S A P B I und 1 2 3 4 5 6 7 8 sowie Betrieb " * 10, "pdf") if "Sperrung" in m], [])


class PeriodTests(unittest.TestCase):
    def test_periods_are_read_in_order_and_a_date_is_no_phone_number(self):
        self.assertEqual(ats.periods(CV), ["03/2022–heute", "09/2017–02/2022", "10/2009–09/2013"])
        self.assertEqual(ats.phone_numbers("10/2009 – 09/2013 und 2019 – 2022 oder 03/2019 - 05/2022"), set())
        self.assertEqual(ats.phone_numbers("Tel +49 30 1234567"), {"0301234567"})

    def test_a_lost_changed_or_reordered_period_is_a_blocker(self):
        text = "\n".join(t for _, _, t in CV_LINES)
        self.assertEqual(ats.station_notes(CV, text), [])
        lost = text.replace("09/2017 – 02/2022", "")
        self.assertEqual([g for _, g, _ in ats.station_notes(CV, lost)], ["Blocker"])
        order = text.replace("09/2017 – 02/2022", "XX").replace("03/2022 – heute", "09/2017 – 02/2022").replace("XX", "03/2022 – heute")
        self.assertEqual([g for _, g, m in ats.station_notes(CV, order)][:1], ["Blocker"])
        self.assertIn("Reihenfolge", ats.station_notes(CV, order)[0][2])

    def test_position_and_employer_must_stand_next_to_their_period(self):
        lines = [t for _, _, t in CV_LINES]
        far = [t for t in lines if t not in ("Softwareentwicklerin", "Beispiel AG, Hamburg")] + ["a", "b", "c", "d", "e", "f", "Softwareentwicklerin", "Beispiel AG, Hamburg"]
        found = ats.station_notes(CV, "\n".join(far))
        self.assertTrue(any("09/2017" in m for _, _, m in found), found)

    def test_markdown_marks_and_typographic_quotes_do_not_break_the_station_match(self):
        body = '### 03/2022 – heute | Leitung „Plattform“ | **Beispiel** AG &amp; Söhne, Berlin\n'
        text = "03/2022 – heute\nLeitung \"Plattform\"\nBeispiel AG & Söhne, Berlin"
        self.assertEqual(ats.station_notes(body, text), [])

    def test_mixed_date_formats_are_a_warning(self):
        mixed = CV.replace("10/2009 – 09/2013", "2009 – 2013")
        text = "\n".join(t for _, _, t in CV_LINES).replace("10/2009 – 09/2013", "2009 – 2013")
        self.assertIn("Uneinheitliche", " ".join(m for _, _, m in ats.station_notes(mixed, text)))


class ContactTests(unittest.TestCase):
    def test_lost_address_and_number_are_blockers_and_never_quoted(self):
        text = "\n".join(t for _, _, t in CV_LINES)
        self.assertEqual(ats.contact_notes(CV, text, "cv", ""), [])
        lost = ats.contact_notes(CV, text.replace("maria.beispiel@beispiel-firma.example", "").replace("+49 30 1234567", ""), "cv", "")
        self.assertEqual([g for _, g, _ in lost], ["Blocker", "Blocker"])
        self.assertNotIn("beispiel-firma", " ".join(m for _, _, m in lost))
        self.assertNotIn("1234567", " ".join(m for _, _, m in lost))

    def test_an_address_that_stands_only_in_the_header_is_named_as_such(self):
        text = "\n".join(t for _, _, t in CV_LINES).replace("maria.beispiel@beispiel-firma.example", "")
        found = ats.contact_notes(CV, text, "cv", "maria.beispiel@beispiel-firma.example")
        self.assertIn("Kopf- oder Fußzeile", found[0][2])

    def test_a_wrapped_address_is_found_and_an_optional_phone_may_be_absent(self):
        text = "Kontakt maria.beispiel@\nbeispiel-firma.example sonst nichts"
        self.assertEqual(ats.contact_notes("maria.beispiel@beispiel-firma.example", text, "letter", ""), [])
        self.assertEqual(ats.contact_notes("nur Text ohne Kontakt", "nur Text ohne Kontakt", "letter", ""), [])

    def test_a_cv_without_an_address_is_a_warning(self):
        self.assertEqual([g for _, g, _ in ats.contact_notes("nur Text", "nur Text", "cv", "")], ["Warnung"])


class LayoutTests(unittest.TestCase):
    def test_a_rating_shown_only_as_symbols_needs_a_review(self):
        gefunden = ats.layout_notes("pdf", None, "", STRUCTURE, "- Python: ★★★★☆\n")
        self.assertEqual([g for _, g, _ in gefunden], ["Prüfung"])
        self.assertEqual(ats.layout_notes("pdf", None, "", STRUCTURE, "- Python: sehr gut\n"), [])

    def test_tables_images_and_columns_in_docx_are_warnings(self):
        xml = '<w:document><w:txbxContent/><w:cols w:num="2"/></w:document>'
        gefunden = ats.layout_notes("docx", xml, "Seite 1", {"tabellen": 1, "bilder": 2}, "")
        self.assertEqual(len(gefunden), 5)
        self.assertEqual({g for _, g, _ in gefunden}, {"Warnung"})

    def test_columns_need_a_gap_at_the_same_place_in_many_lines(self):
        zeilen = lambda n: [(x, y * 12.0, x + 80, y * 12.0 + 9) for y in range(n) for x in (72, 330)]
        self.assertTrue(ats.has_columns(595, zeilen(12)))
        single = [(72, y * 12.0, 400, y * 12.0 + 9) for y in range(12)]
        self.assertFalse(ats.has_columns(595, single))
        stations = single + [(x, 200 + y * 40.0, x + 60, 209 + y * 40.0) for y in range(2) for x in (72, 330)]
        self.assertFalse(ats.has_columns(595, stations))  # a period column in a few lines is no page layout


class SectionTests(unittest.TestCase):
    def test_headings_must_stand_on_their_own_line_and_keep_their_content(self):
        text = "\n".join(t for _, _, t in CV_LINES)
        self.assertEqual(ats.section_notes(CV, text, "pdf"), [])
        merged = text.replace("Berufserfahrung\n", "Berufserfahrung ")
        self.assertIn("eigene Zeile", ats.section_notes(CV, merged, "pdf")[0][2])
        moved = "\n".join(t for _, _, t in CV_LINES if t != "Ausbildung und Studium")
        self.assertTrue(any("Ausbildung" in m for _, _, m in ats.section_notes(CV, moved, "pdf")))
        content_gone = "\n".join(t for _, _, t in CV_LINES[:11]) + "\nAusbildung und Studium\n" + "\n".join(t for _, _, t in CV_LINES[11:13])
        self.assertTrue(any("steht nicht unter seiner Überschrift" in m for _, _, m in ats.section_notes(CV, content_gone, "pdf")))

    def test_other_headings_are_no_problem_and_html_is_not_judged_by_lines(self):
        text = "Werdegang\n" + "\n".join(t for _, _, t in CV_LINES[3:])
        self.assertEqual(ats.section_notes("## Werdegang\n- Kubernetes betrieben und Plattform gesenkt\n", "Werdegang\nKubernetes betrieben und Plattform gesenkt", "pdf"), [])
        self.assertEqual(ats.section_notes(CV, text, "html"), [])


class FontTests(unittest.TestCase):
    LISTING = """name                                 type              encoding         emb sub uni object ID
------------------------------------ ----------------- ---------------- --- --- --- ---------
ABCDEF+LatoRegular                   CID TrueType      Identity-H       yes yes yes     12  0
Bitmap                               Type 3            Custom           yes no  no      13  0
Fantasie                             TrueType          WinAnsi          no  no  no      14  0
Helvetica                            Type 1            WinAnsi          no  no  no      15  0
"""

    def test_type3_and_not_embedded_fonts_are_found_and_the_standard_fonts_are_not(self):
        gefunden = ats.font_notes(self.LISTING)
        self.assertEqual([(g, m[:1]) for _, g, m in gefunden], [("Warnung", "1"), ("Warnung", "1")])
        self.assertIn("Typ 3", gefunden[0][2])
        self.assertIn("nicht eingebettet", gefunden[1][2])
        self.assertEqual(ats.font_notes("\n".join(self.LISTING.splitlines()[:3])), [])


class KeyTermTests(unittest.TestCase):
    def test_short_terms_are_found_and_a_prefix_of_another_word_is_not_a_hit(self):
        self.assertTrue(ats.term_in("SQL", "erfahrung mit sql und python"))
        self.assertTrue(ats.term_in("CI/CD", "ci/cd pipelines gebaut"))
        self.assertFalse(ats.term_in("Java", "javascript und typescript"))
        self.assertTrue(ats.term_in("Java", "java und kotlin"))
        self.assertTrue(ats.term_in("Führung", "führungsrolle im team"))
        self.assertFalse(ats.term_in("", "text"))

    ANALYSE = """# Analyse

| Anforderung | Priorität | Kompetenz | Nachweis | Stärke | Status | Verwendung | Offen |
|---|---|---|---|---|---|---|---|
| „Kubernetes“ | Muss | Kubernetes / Plattformbetrieb | Station | stark | erfüllt | Lebenslauf | — |
| „Terraform“ | Muss | Terraform | Station | mittel | erfüllt | Lebenslauf | — |
| „Budgetverantwortung“ | Wunsch | Budget | — | — | nicht belegt | Rückfrage | — |
"""

    def test_an_evidenced_term_the_cv_does_not_name_is_a_warning(self):
        gefunden = ats.key_terms(self.ANALYSE, "- Kubernetes betrieben, Plattformbetrieb verantwortet\n")
        self.assertEqual([(g, "Terraform" in m) for _, g, m in gefunden], [("Warnung", True)])

    def test_a_term_of_the_posting_that_the_analysis_does_not_evidence_needs_a_review(self):
        gefunden = ats.key_terms(self.ANALYSE, "- Kubernetes, Plattformbetrieb, Terraform und Budget verantwortet\n")
        self.assertEqual([(g, "Budget" in m) for _, g, m in gefunden], [("Prüfung", True)])

    def test_the_toolchain_inserts_nothing(self):
        text = "- Kubernetes betrieben\n"
        ats.key_terms(self.ANALYSE, text)
        self.assertEqual(text, "- Kubernetes betrieben\n")


@unittest.skipUnless(POPPLER, "poppler-utils (pdftotext, pdfinfo, pdffonts, pdfimages) is not installed")
class PdfTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = self._tmp.name

    def pruefe(self, pdf, body=CV, doc_type="cv", ohne_werkzeuge=False):
        pfad = write(self.tmp, "x.pdf", pdf)
        text = qa.extrahiere(pfad, "pdf")[0] or ""
        text = qa.ex.strip_page_numbers(text)
        if ohne_werkzeuge:
            with unittest.mock.patch.object(ats.shutil, "which", return_value=None):
                return ats.check_file(pfad, "pdf", doc_type, body, text, STRUCTURE)
        return ats.check_file(pfad, "pdf", doc_type, body, text, STRUCTURE)

    def test_a_clean_pdf_has_no_finding(self):
        gefunden, luecken = self.pruefe(positioned_pdf([CV_LINES]))
        self.assertEqual((gefunden, luecken), ([], {}))

    def test_a_file_that_is_no_pdf_or_is_cut_off_is_a_blocker(self):
        gefunden, _ = self.pruefe(b"%PDF-1.4\n1 0 obj\n<<>>\n")
        self.assertEqual([(c, g) for c, g, _ in gefunden if c == "ATS-01"], [("ATS-01", "Blocker")])
        gefunden, _ = self.pruefe(b"das ist ein Text")
        self.assertIn("Blocker", [g for c, g, _ in gefunden if c == "ATS-01"])

    def test_an_image_text_outside_the_page_and_columns_are_warnings(self):
        gefunden, _ = self.pruefe(positioned_pdf([CV_LINES], image=True))
        self.assertTrue(any(c == "ATS-05" and "Bild" in m for c, _, m in gefunden), gefunden)
        out = CV_LINES[:-1] + [(560, 536, CV_LINES[-1][2])]
        gefunden, _ = self.pruefe(positioned_pdf([out]))
        self.assertTrue(any(c == "ATS-05" and "außerhalb der Seite" in m for c, _, m in gefunden), gefunden)
        zweispaltig = [(x, 780 - 14 * i, f"Kubernetes Plattform betrieben {i}") for i in range(14) for x in (72, 330)]
        gefunden, _ = self.pruefe(positioned_pdf([zweispaltig]), body="## Profil\n- Kubernetes\n")
        self.assertTrue(any(c == "ATS-05" and "Mehrspaltig" in m for c, _, m in gefunden), gefunden)

    def test_a_font_that_is_not_embedded_is_a_warning(self):
        gefunden, _ = self.pruefe(positioned_pdf([CV_LINES], font="/Type /Font /Subtype /TrueType /BaseFont /Fantasie"))
        self.assertTrue(any(c == "ATS-01" and "nicht eingebettet" in m for c, _, m in gefunden), gefunden)

    def test_missing_tools_are_gaps_and_never_a_pass(self):
        gefunden, luecken = self.pruefe(positioned_pdf([CV_LINES]), ohne_werkzeuge=True)
        self.assertEqual(sorted(luecken), ["ATS-01", "ATS-05"])


@unittest.skipUnless(docx is not None, "python-docx is not installed")
class DocxTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.pfad = os.path.join(self._tmp.name, "x.docx")

    def baue(self, aendere=None):
        dokument = docx.Document()
        dokument.add_heading("Lebenslauf", 1)
        dokument.add_heading("Persönliche Daten", 2)
        dokument.add_paragraph("Maria Beispiel, Berlin, maria.beispiel@beispiel-firma.example, +49 30 1234567")
        dokument.add_heading("Berufserfahrung", 2)
        for _, _, zeile in CV_LINES[4:13]:
            dokument.add_paragraph(zeile)
        dokument.add_heading("Ausbildung und Studium", 2)
        dokument.add_paragraph(CV_LINES[-1][2])
        if aendere:
            aendere(dokument)
        dokument.save(self.pfad)

    def pruefe(self, body=CV):
        text = qa.extrahiere(self.pfad, "docx")
        return ats.check_file(self.pfad, "docx", "cv", body, text[0], text[1])[0]

    def test_a_clean_docx_has_no_finding(self):
        self.baue()
        self.assertEqual(self.pruefe(), [])

    def test_hidden_white_and_tiny_text_is_a_blocker(self):
        def verstecke(dokument):
            absatz = dokument.add_paragraph()
            absatz.add_run("Kubernetes Terraform").font.hidden = True
            weiss = absatz.add_run("Python")
            weiss.font.color.rgb = docx.shared.RGBColor(255, 255, 255)
            absatz.add_run("Go").font.size = docx.shared.Pt(1)
        self.baue(verstecke)
        gefunden = self.pruefe()
        self.assertTrue(any(c == "ATS-01" and g == "Blocker" and "1 ausgeblendet, 1 weiß, 1 winzig" in m for c, g, m in gefunden), gefunden)

    def test_white_text_on_a_dark_band_is_a_warning_not_a_blocker(self):
        def band(dokument):
            absatz = dokument.add_paragraph()
            absatz.add_run("Kontakt").font.color.rgb = docx.shared.RGBColor(255, 255, 255)
            absatz._p.get_or_add_pPr().append(docx.oxml.parse_xml('<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:fill="1F3864"/>'))
        self.baue(band)
        gefunden = [(g, m) for c, g, m in self.pruefe() if c == "ATS-01"]
        self.assertEqual([g for g, _ in gefunden], ["Warnung"], gefunden)

    def test_a_title_block_counts_as_the_level_one_heading(self):
        xml = '<w:p><w:pStyle w:val="Title"/></w:p><w:p><w:pStyle w:val="Heading2"/></w:p>'
        self.assertEqual(ats.docx_heading_notes(xml, "# Name\n\n## Berufserfahrung\n"), [])
        self.assertEqual(ats.docx_heading_notes(xml, "# Name\n\n```\n# kein Titel\n```\n\n## Berufserfahrung\n"), [])

    def test_a_table_header_text_and_a_contact_that_only_stands_in_the_header_are_found(self):
        def aendere(dokument):
            dokument.add_table(rows=1, cols=2)
            dokument.sections[0].header.paragraphs[0].text = "maria.beispiel@beispiel-firma.example"
            for absatz in dokument.paragraphs:
                if "maria.beispiel" in absatz.text:
                    absatz.text = "Maria Beispiel, Berlin"
        self.baue(aendere)
        gefunden = self.pruefe()
        self.assertTrue(any(c == "ATS-05" and "Tabelle" in m for c, _, m in gefunden), gefunden)
        self.assertTrue(any(c == "ATS-05" and "Kopf- oder Fußzeile" in m for c, _, m in gefunden), gefunden)
        self.assertTrue(any(c == "ATS-04" and "nur in der Kopf" in m for c, _, m in gefunden), gefunden)

    def test_a_heading_without_its_style_is_found(self):
        def flach(dokument):
            for absatz in dokument.paragraphs:
                if absatz.text == "Berufserfahrung":
                    absatz.style = dokument.styles["Normal"]
        self.baue(flach)
        gefunden = self.pruefe()
        self.assertTrue(any(c == "ATS-02" and "Formatvorlagen" in m for c, _, m in gefunden), gefunden)

    def test_a_broken_archive_and_a_broken_xml_are_blockers(self):
        self.baue()
        with open(self.pfad, "rb") as handle:
            daten = handle.read()
        with open(self.pfad, "wb") as handle:
            handle.write(daten[:len(daten) // 2])
        self.assertEqual(ats.docx_parts(self.pfad)[0], None)
        self.baue()
        with zipfile.ZipFile(self.pfad) as quelle, zipfile.ZipFile(self.pfad + ".neu", "w") as ziel:
            for name in quelle.namelist():
                inhalt = quelle.read(name)
                ziel.writestr(name, inhalt[:-20] if name == "word/document.xml" else inhalt)
        shutil.move(self.pfad + ".neu", self.pfad)
        xml, _, gefunden = ats.docx_parts(self.pfad)
        self.assertEqual((xml, [g for _, g, _ in gefunden]), (None, ["Blocker"]))


class ReportTests(tq.Base):
    """The checks inside the one report."""

    def exportiere(self, name, fmt, doc_type):
        lauf = subprocess.run([sys.executable, tq.EXPORTER, os.path.join(self.ordner, name), "--format", fmt, "--doc-type", doc_type,
                               "--out-dir", os.path.join(self.ordner, "export")], capture_output=True, text=True)
        self.assertEqual(lauf.returncode, 0, lauf.stderr)

    def ats_status(self, bericht):
        return {c["id"]: c["status"] for c in bericht["checks"] if c["id"].startswith("ATS-") or c["id"] in ("EX-03", "KW-01")}

    def test_without_an_export_the_ats_checks_are_not_applicable_and_nothing_is_validated(self):
        bericht = self.lauf()
        self.assertEqual({v for k, v in self.ats_status(bericht).items() if k != "KW-01"}, {qa.NA})
        self.assertEqual((bericht["ats"]["umfang"], bericht["ats"]["technisch_validiert"]), ("nicht_geprueft", False))
        self.assertIn("keine Zusage", bericht["ats"]["hinweis"])

    @unittest.skipUnless(tq.HAS_PANDOC, "pandoc is not installed")
    def test_a_markdown_copy_is_not_an_ats_export(self):
        self.exportiere("lebenslauf.md", "md", "cv")
        bericht = self.lauf()
        self.assertEqual(self.check(bericht, "ATS-03")["status"], qa.NA)
        self.assertFalse(bericht["ats"]["technisch_validiert"])

    @unittest.skipUnless(tq.HAS_PANDOC, "pandoc is not installed")
    def test_the_html_export_of_the_exporter_passes_the_ats_checks(self):
        self.exportiere("lebenslauf.md", "html", "cv")
        bericht = self.lauf()
        for check_id, status in self.ats_status(bericht).items():
            self.assertIn(status, (qa.PASS, qa.WARNING, qa.NA), (check_id, self.check(bericht, check_id)["befunde"]))
        self.assertEqual(bericht["ats"]["dateien"], ["export/lebenslauf.html"])

    @unittest.skipUnless(tq.HAS_PANDOC and docx is not None, "pandoc or python-docx is not installed")
    def test_the_docx_export_of_the_exporter_passes_the_ats_checks(self):
        self.exportiere("lebenslauf.md", "docx", "cv")
        bericht = self.lauf()
        for check_id in ("EX-02", "EX-03", "ATS-01", "ATS-02", "ATS-03", "ATS-04", "ATS-05"):
            self.assertNotEqual(self.check(bericht, check_id)["status"], qa.FAIL, (check_id, self.check(bericht, check_id)["befunde"]))
        self.assertTrue(bericht["ats"]["technisch_validiert"], bericht["ats"])

    @unittest.skipUnless(tq.HAS_PANDOC and docx is not None, "pandoc or python-docx is not installed")
    def test_a_defect_in_the_export_blocks_the_package_and_a_changed_export_is_checked_again(self):
        self.exportiere("lebenslauf.md", "docx", "cv")
        pfad = os.path.join(self.ordner, "export", "lebenslauf.docx")
        with zipfile.ZipFile(pfad) as quelle, zipfile.ZipFile(pfad + ".neu", "w") as ziel:
            for name in quelle.namelist():
                inhalt = quelle.read(name)
                if name == "word/document.xml":
                    inhalt = inhalt.replace(b"<w:r>", b"<w:r><w:rPr><w:vanish/></w:rPr>", 1)
                ziel.writestr(name, inhalt)
        shutil.move(pfad + ".neu", pfad)
        manifest_pfad = pfad + ".manifest.json"
        manifest = json.loads(tq.read_file(manifest_pfad))
        manifest["ausgabe"]["sha256"] = qa.datei_hash(pfad)  # a consistent manifest: only the content differs
        with open(manifest_pfad, "w", encoding="utf-8") as handle:
            json.dump(manifest, handle)
        bericht = self.lauf()
        self.assertEqual(self.check(bericht, "ATS-01")["status"], qa.FAIL)
        self.assertEqual((bericht["gesamt"], bericht["freigabefaehig"], bericht["ats"]["technisch_validiert"]), ("BLOCKIERT", False, False))
        self.assertTrue(all("maria" not in b["befund"].lower() for c in bericht["checks"] for b in c["befunde"]))

    def test_every_check_has_a_dimension_and_the_report_has_no_total_score(self):
        bericht = self.lauf()
        self.assertTrue(all(c["dimension"] == qa.DIMENSION[c["id"]] for c in bericht["checks"]))
        self.assertEqual(set(qa.DIMENSION), set(qa.KATALOG))
        self.assertEqual(set(bericht["dimensionen"]), {"technisch", "evidenz", "formal", "schluesselbegriffe", "passung"})
        self.assertFalse([k for k in bericht if "score" in k.lower() or "punkte" in k.lower()])
        self.assertEqual(self.check(bericht, "KW-01")["dimension"], "schluesselbegriffe")

    def test_the_key_terms_of_the_analysis_are_checked_on_the_cv(self):
        basis = self.check(self.lauf(), "KW-01")  # the fixture names no "Personalführung": an evidenced term the CV does not state
        self.assertEqual((basis["status"], len(basis["befunde"])), (qa.WARNING, 1))
        self.aendere("analyse.md", "Wunsch | Budget | — | — | nicht belegt", "Wunsch | Kubernetes | — | — | nicht belegt")
        status = self.check(self.lauf(), "KW-01")
        self.assertEqual(status["status"], qa.REVIEW, status["befunde"])

    def test_without_the_analysis_the_key_terms_are_not_applicable(self):
        os.remove(os.path.join(self.ordner, "analyse.md"))
        self.assertEqual(self.check(self.lauf(), "KW-01")["status"], qa.NA)


class SafetyTests(unittest.TestCase):
    def test_the_library_sends_nothing_and_runs_only_the_poppler_tools(self):
        quelle = tq.read_file(os.path.join(ROOT, "scripts", "ats_checks.py"))
        imports = set(re.findall(r"^(?:import|from) (\w+)", quelle, re.M))
        self.assertFalse(imports & {"socket", "smtplib", "urllib", "http", "requests", "ssl", "ftplib", "imaplib", "webbrowser", "workflow"}, imports)
        self.assertEqual(sorted(set(re.findall(r"subprocess\.run\(\[\"(\w+)\"", quelle))), ["pdffonts", "pdfimages", "pdfinfo", "pdftotext"])
        self.assertNotIn("open(path, \"w", quelle)


if __name__ == "__main__":
    unittest.main()
